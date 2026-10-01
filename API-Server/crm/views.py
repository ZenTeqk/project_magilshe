import csv
import io
import uuid
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.db import models, transaction
from django.db.models import Count, ProtectedError, Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import BasePermission, IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from cakery.common import aws_domain_name, get_s3_session
from cakery.file_upload_to_aws import put_image_to_s3
from customer.rbac import effective_permissions, is_super_admin
from .models import (
    CRMDelivery, CRMDeliveryDamageImage, CRMInvoice, CRMOption, CRMOrder, CRMOrderItem, CustomerCommunicationLog, CustomerRecurringEvent,
    Lead, LeadActivity, LeadReferenceImage, Quotation, QuotationActivity, QuotationItem, QuotationPayment,
    pricing_quantity,
)
from .customer_matching import find_existing_customer, sync_customer_from_crm_details
from .agenda import MAX_RANGE_DAYS, build_agenda
from .actionable import (
    DELIVERY_ACTIONABLE_STATUSES, INVOICE_ACTIONABLE_STATUSES,
    LEAD_ACTIONABLE_STATUSES, ORDER_ACTIONABLE_STATUSES,
    QUOTATION_ACTIONABLE_STATUSES, status_payload,
)
from .recurring_events import recurring_event_summary, sync_recurring_event
from .uploads import validate_image_upload
from whatsapp.services import trigger_event as whatsapp_trigger_event
from .serializers import (
    CRMDeliveryDamageImageSerializer, CRMDeliverySerializer, CRMInvoiceSerializer, CRMOptionSerializer, CRMOrderSerializer,
    CustomerRecurringEventSerializer, LeadReferenceImageSerializer, LeadSerializer,
    QuotationListSerializer, QuotationPaymentSerializer, QuotationSerializer, CRMCustomerUpdateSerializer,
    create_quotation_version, is_quotation_locked,
)


def _wants_archived(request):
    return str(request.query_params.get("archived") or "").lower() in {"1", "true", "yes"}


def _filter_archived(queryset, request):
    if _wants_archived(request):
        return queryset.filter(is_archived=True) if is_super_admin(request.user) else queryset.none()
    return queryset.filter(is_archived=False)


def _archive_record(instance, user, label):
    instance.is_archived = True
    instance.archived_at = timezone.now()
    instance.archived_by = user
    instance.save(update_fields=("is_archived", "archived_at", "archived_by"))
    return Response({"message": f"{label} archived."}, status=status.HTTP_200_OK)


def _restore_record(instance, label):
    instance.is_archived = False
    instance.archived_at = None
    instance.archived_by = None
    instance.save(update_fields=("is_archived", "archived_at", "archived_by"))
    return Response({"message": f"{label} restored."}, status=status.HTTP_200_OK)


def _require_super_admin(request):
    if not is_super_admin(request.user):
        return Response({"detail": "Only Super Admin can perform this action."}, status=status.HTTP_403_FORBIDDEN)
    return None


def _event_fields(instance, *fields):
    """Snapshot of the fields that decide a record's customer event."""
    return tuple(getattr(instance, field) for field in fields)


class ArchivedLifecycleMixin:
    lifecycle_label = "Record"

    def get_queryset(self):
        if getattr(self, "action", None) in {"restore", "permanent_delete"}:
            return super().get_queryset()
        if getattr(self, "action", None) in {"retrieve", "update", "partial_update"} and is_super_admin(self.request.user):
            return super().get_queryset()
        return _filter_archived(super().get_queryset(), self.request)

    def destroy(self, request, *args, **kwargs):
        return _archive_record(self.get_object(), request.user, self.lifecycle_label)

    @action(detail=True, methods=["post"], url_path="restore")
    def restore(self, request, pk=None):
        denied = _require_super_admin(request)
        if denied:
            return denied
        return _restore_record(self.get_object(), self.lifecycle_label)

    @action(detail=True, methods=["delete"], url_path="permanent-delete")
    def permanent_delete(self, request, pk=None):
        denied = _require_super_admin(request)
        if denied:
            return denied
        instance = self.get_object()
        blocker = self.permanent_delete_blocker(instance)
        if blocker:
            return Response(
                {"message": blocker, "dependencies": self.deletion_dependencies(instance)},
                status=status.HTTP_409_CONFLICT,
            )
        try:
            with transaction.atomic():
                self.prepare_permanent_delete(instance)
                instance.delete()
        except ProtectedError as exc:
            dependencies = self.protected_error_dependencies(exc)
            return Response(
                {
                    "message": f"{self.lifecycle_label} cannot be permanently deleted because it has protected related records. Delete linked records first.",
                    "dependencies": dependencies,
                },
                status=status.HTTP_409_CONFLICT,
            )
        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(detail=True, methods=["get"], url_path="deletion-dependencies")
    def deletion_dependencies_action(self, request, pk=None):
        denied = _require_super_admin(request)
        if denied:
            return denied
        return Response(self.deletion_dependencies(self.get_object()))

    def permanent_delete_blocker(self, instance):
        return None

    def prepare_permanent_delete(self, instance):
        return None

    def deletion_dependencies(self, instance):
        blocker = self.permanent_delete_blocker(instance)
        if not blocker:
            return {"has_dependencies": False, "links": []}
        return {
            "has_dependencies": True,
            "links": [{
                "label": "Blocked",
                "count": 1,
                "records": [{"id": instance.pk, "message": blocker}],
            }],
        }

    def protected_error_dependencies(self, exc):
        protected = list(getattr(exc, "protected_objects", []) or [])
        records = []
        for obj in protected[:20]:
            record = {
                "id": getattr(obj, "pk", None),
                "message": str(obj),
            }
            if isinstance(obj, CRMOrder):
                record.update({
                    "order_no": obj.order_no,
                    "customer_name": obj.customer_name,
                    "status": obj.status,
                    "record_type": "orders",
                })
            elif isinstance(obj, CRMInvoice):
                record.update({
                    "invoice_no": obj.invoice_no,
                    "customer_name": obj.customer_name,
                    "status": obj.status,
                    "record_type": "invoices",
                })
            elif isinstance(obj, CRMDelivery):
                record.update({
                    "delivery_no": obj.delivery_no,
                    "customer_name": obj.recipient_name,
                    "status": obj.delivery_status,
                    "record_type": "deliveries",
                })
            elif isinstance(obj, Quotation):
                record.update({
                    "quotation_no": obj.quotation_no,
                    "customer_name": obj.customer_name,
                    "status": obj.status,
                    "record_type": "quotations",
                })
            records.append(record)
        return {
            "has_dependencies": bool(records),
            "links": [{
                "label": "Protected related records",
                "count": len(protected),
                "records": records,
            }] if records else [],
        }


class LeadViewSet(ArchivedLifecycleMixin, viewsets.ModelViewSet):
    lifecycle_label = "Lead"
    serializer_class = LeadSerializer
    queryset = Lead.objects.select_related("cake", "assigned_to").all()

    def get_queryset(self):
        qs = super().get_queryset()
        p = self.request.query_params
        if p.get("search"):
            term = p["search"]
            qs = qs.filter(Q(lead_no__icontains=term) | Q(customer_name__icontains=term) | Q(phone__icontains=term))
        if p.get("action_required", "").lower() in ("1", "true", "yes"):
            qs = qs.filter(status__in=LEAD_ACTIONABLE_STATUSES)
        for field in ("status", "occasion", "cake_type", "event_date"):
            if p.get(field):
                qs = qs.filter(**{field: p[field]})
        if p.get("created_date"):
            qs = qs.filter(created_on__date=p["created_date"])
        if p.get("from_date"):
            qs = qs.filter(created_on__date__gte=p["from_date"])
        if p.get("to_date"):
            qs = qs.filter(created_on__date__lte=p["to_date"])
        return qs

    def permanent_delete_blocker(self, instance):
        if Quotation.objects.filter(lead=instance).exists():
            return "Lead cannot be permanently deleted because it has a related quotation. Archive it instead."
        return None

    def deletion_dependencies(self, instance):
        links = []
        quotations = Quotation.objects.filter(lead=instance).order_by("-id")
        if quotations.exists():
            links.append({
                "label": "Quotations",
                "count": quotations.count(),
                "records": [{
                    "id": quotation.id,
                    "quotation_no": quotation.quotation_no,
                    "customer_name": quotation.customer_name,
                    "status": quotation.status,
                    "record_type": "quotations",
                } for quotation in quotations[:20]],
            })
        return {"has_dependencies": bool(links), "links": links}

    @action(detail=False, methods=["get"], url_path="crm-dashboard")
    def crm_dashboard(self, request):
        """One compact, cacheable snapshot avoids loading every module list."""
        from_date = request.query_params.get("from_date")
        to_date = request.query_params.get("to_date")

        def in_range(queryset, field):
            if from_date:
                queryset = queryset.filter(**{f"{field}__date__gte": from_date})
            if to_date:
                queryset = queryset.filter(**{f"{field}__date__lte": to_date})
            return queryset

        leads_qs = in_range(Lead.objects.filter(is_archived=False), "created_on")
        quotations_qs = in_range(Quotation.objects.filter(is_archived=False), "created_on")
        orders_qs = in_range(CRMOrder.objects.filter(is_archived=False), "confirmed_on")
        invoices_qs = in_range(CRMInvoice.objects.filter(is_archived=False), "generated_on")
        deliveries_qs = in_range(CRMDelivery.objects.filter(is_archived=False), "created_on")
        follow_ups_qs = leads_qs.filter(status="Follow-up", follow_up_date__isnull=False)
        lead_counts = dict(leads_qs.values_list("status").annotate(total=models.Count("id")))
        status_counts = {
            "quotations": dict(quotations_qs.values_list("status").annotate(total=Count("id"))),
            "orders": dict(orders_qs.values_list("status").annotate(total=Count("id"))),
            "invoices": dict(invoices_qs.values_list("status").annotate(total=Count("id"))),
            "deliveries": dict(deliveries_qs.values_list("delivery_status").annotate(total=Count("id"))),
        }
        actionable = {
            "leads": leads_qs.filter(status__in=LEAD_ACTIONABLE_STATUSES).count(),
            "follow_ups": follow_ups_qs.count(),
            "quotations": quotations_qs.filter(status__in=QUOTATION_ACTIONABLE_STATUSES).count(),
            "orders": orders_qs.filter(status__in=ORDER_ACTIONABLE_STATUSES).count(),
            "invoices": invoices_qs.filter(status__in=INVOICE_ACTIONABLE_STATUSES, balance_amount__gt=0).count(),
            "deliveries": deliveries_qs.filter(delivery_status__in=DELIVERY_ACTIONABLE_STATUSES).count(),
        }
        document_counts = {
            "quotations": quotations_qs.count(), "orders": orders_qs.count(),
            "invoices": invoices_qs.count(), "deliveries": deliveries_qs.count(),
        }
        funnel = {
            "total_leads": leads_qs.count(),
            "lead_only": leads_qs.filter(quotation__isnull=True).count(),
            "moved_to_quotation": quotations_qs.count(),
            "moved_to_order": orders_qs.count(),
            "moved_to_invoice": invoices_qs.count(),
            "moved_to_delivery": deliveries_qs.count(),
            "completed": leads_qs.filter(status="Completed").count(),
            "won_not_quoted": leads_qs.filter(status="Won", quotation__isnull=True).count(),
            "lost": leads_qs.filter(status="Lost").count(),
        }
        workflows = []
        quotations = in_range(Quotation.objects.filter(is_archived=False).select_related(
            "lead", "crm_order", "crm_order__invoice", "crm_order__delivery"
        ), "updated_on").order_by("-updated_on")[:10]
        for quotation in quotations:
            order = getattr(quotation, "crm_order", None)
            invoice = getattr(order, "invoice", None) if order else None
            delivery = getattr(order, "delivery", None) if order else None
            if delivery:
                stage, current_status, target_id = "Delivery", delivery.delivery_status, delivery.id
            elif invoice:
                stage, current_status, target_id = "Invoice", invoice.status, invoice.id
            elif order:
                stage, current_status, target_id = "Order", order.status, order.id
            else:
                stage, current_status, target_id = "Quotation", quotation.status, quotation.id
            workflows.append({
                "customer_name": quotation.customer_name,
                "lead_no": quotation.lead.lead_no if quotation.lead else None,
                "quotation_no": quotation.quotation_no,
                "stage": stage, "status": current_status, "target_id": target_id,
                "updated_on": quotation.updated_on,
            })
        payload = {"leads": {"total": sum(lead_counts.values()), **lead_counts},
                   "documents": document_counts, "status_counts": status_counts,
                   "actionable": actionable, "actionable_statuses": status_payload(),
                   "funnel": funnel, "workflows": workflows}
        return Response(payload)


    def perform_create(self, serializer):
        customer = serializer.validated_data.get("customer")
        customer = sync_customer_from_crm_details(
            {**serializer.validated_data, "customer": customer}, self.request.user
        )
        lead = serializer.save(created_by=self.request.user, customer=customer)
        LeadActivity.objects.create(lead=lead, action="Lead Created", staff=self.request.user)
        sync_recurring_event(lead)

    def perform_update(self, serializer):
        previous_status = serializer.instance.status
        previous_event = _event_fields(serializer.instance, "occasion", "event_date", "customer_id")
        previous_notes = serializer.instance.event_notes
        save_kwargs = {"updated_by": self.request.user}
        if "customer" in serializer.validated_data or "phone" in serializer.validated_data:
            synced_customer = sync_customer_from_crm_details(
                {
                    **serializer.validated_data,
                    "customer": serializer.validated_data.get("customer", serializer.instance.customer),
                    "customer_name": serializer.validated_data.get("customer_name", serializer.instance.customer_name),
                    "phone": serializer.validated_data.get("phone", serializer.instance.phone),
                    "email": serializer.validated_data.get("email", serializer.instance.email),
                    "address": serializer.validated_data.get("address", serializer.instance.address),
                    "city": serializer.validated_data.get("city", serializer.instance.city),
                    "pincode": serializer.validated_data.get("pincode", serializer.instance.pincode),
                },
                self.request.user,
            )
            if synced_customer is not None:
                save_kwargs["customer"] = synced_customer
        lead = serializer.save(**save_kwargs)
        LeadActivity.objects.create(lead=lead, action="Lead Edited", staff=self.request.user)
        if _event_fields(lead, "occasion", "event_date", "customer_id") != previous_event:
            sync_recurring_event(lead)
        elif lead.event_notes != previous_notes:
            sync_recurring_event(lead, create=False)
        if previous_status != lead.status:
            LeadActivity.objects.create(
                lead=lead, action="Status Changed", staff=self.request.user,
                description=f"{previous_status} → {lead.status}",
            )
        if lead.status == "Follow-up" and (
            previous_status != lead.status or "follow_up_date" in serializer.validated_data
        ):
            LeadActivity.objects.create(
                lead=lead, action="Follow-up Scheduled", staff=self.request.user,
                description=f"{lead.follow_up_date} {lead.follow_up_time or ''} — {lead.follow_up_notes}",
            )

    @action(detail=False, methods=["get"], url_path="customer-match")
    def customer_match(self, request):
        customer = find_existing_customer(request.query_params.get("phone"))
        if not customer:
            return Response({"data": None})
        addresses = customer.customeraddress_set.order_by("-is_primary", "id")
        return Response({"data": {
            "id": customer.id,
            "name": customer.first_name,
            "phone": customer.phone,
            "email": customer.email,
            "address_details": [{
                "id": address.id,
                "title": address.title,
                "person_name": address.person_name,
                "phone": address.phone,
                "address": address.address,
                "city": address.city,
                "pincode": address.pincode,
                "is_primary": address.is_primary,
            } for address in addresses],
        }})

    @action(detail=False, methods=["get"], url_path="follow-ups")
    def follow_ups(self, request):
        today = timezone.localdate()
        active = self.get_queryset().filter(
            status="Follow-up", follow_up_date__isnull=False
        ).order_by("follow_up_date", "follow_up_time", "id")
        if request.query_params.get("from_date"):
            active = active.filter(follow_up_date__gte=request.query_params["from_date"])
        if request.query_params.get("to_date"):
            active = active.filter(follow_up_date__lte=request.query_params["to_date"])
        completed_activity = LeadActivity.objects.filter(
            action="Follow-up Completed", created_on__date=today
        )
        completed_leads = Lead.objects.filter(
            activities__in=completed_activity
        ).select_related("cake", "assigned_to").distinct()
        return Response({
            "results": self.get_serializer(active, many=True).data,
            "completed_results": self.get_serializer(completed_leads, many=True).data,
            "current_staff": str(request.user),
            "summary": {
                "today": active.filter(follow_up_date=today).count(),
                "overdue": active.filter(follow_up_date__lt=today).count(),
                "completed_today": completed_activity.count(),
                "assigned_to_me": active.filter(assigned_to=request.user).count(),
            },
        })

    @action(detail=True, methods=["post"], url_path="follow-up-note")
    def follow_up_note(self, request, pk=None):
        lead = self.get_object()
        note = str(request.data.get("note") or "").strip()
        if not note:
            return Response({"note": "A follow-up note is required."}, status=400)
        lead.follow_up_notes = note
        lead.updated_by = request.user
        lead.save(update_fields=("follow_up_notes", "updated_by", "updated_on"))
        activity = LeadActivity.objects.create(
            lead=lead, action="Follow-up Note", staff=request.user, description=note
        )
        return Response({
            "lead": self.get_serializer(lead).data,
            "activity": {
                "id": activity.id, "action": activity.action,
                "description": activity.description, "staff_name": str(request.user),
                "created_on": activity.created_on,
            },
        })

    @action(detail=True, methods=["post"], url_path="complete-follow-up")
    def complete_follow_up(self, request, pk=None):
        lead = self.get_object()
        if lead.status != "Follow-up" or not lead.follow_up_date:
            return Response({"message": "This lead has no active follow-up."}, status=400)
        description = str(request.data.get("note") or "").strip() or "Follow-up completed"
        lead.follow_up_date = None
        lead.follow_up_time = None
        lead.updated_by = request.user
        lead.save(update_fields=("follow_up_date", "follow_up_time", "updated_by", "updated_on"))
        LeadActivity.objects.create(
            lead=lead, action="Follow-up Completed", staff=request.user, description=description
        )
        return Response({"message": "Follow-up completed.", "data": self.get_serializer(lead).data})

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def convert(self, request, pk=None):
        lead = Lead.objects.select_for_update().get(pk=self.get_object().pk)
        if "crm.quotations.create" not in effective_permissions(request.user):
            return Response(
                {"detail": "You do not have permission to create quotations."},
                status=status.HTTP_403_FORBIDDEN,
            )
        mark_won = request.data.get("mark_won") is True
        if mark_won and lead.status != "Won":
            previous_status = lead.status
            lead.status = "Won"
            lead.updated_by = request.user
            lead.save(update_fields=("status", "updated_by", "updated_on"))
            LeadActivity.objects.create(
                lead=lead, action="Status Changed", staff=request.user,
                description=f"{previous_status} to Won · Won & Generate Quotation",
            )
        primary_address = None
        customer_newly_linked = False
        if lead.customer is None:
            lead.customer = sync_customer_from_crm_details({
                "customer_name": lead.customer_name,
                "phone": lead.phone,
                "email": lead.email,
                "address": lead.address,
                "city": lead.city,
                "pincode": lead.pincode,
            }, request.user)
            if lead.customer is not None:
                lead.save(update_fields=("customer", "updated_on"))
                customer_newly_linked = True
        if lead.customer:
            primary_address = lead.customer.customeraddress_set.order_by("-is_primary", "id").first()
        quotation, created = Quotation.objects.get_or_create(
            lead=lead,
            defaults={
                "customer": lead.customer,
                "customer_name": lead.customer_name,
                "phone": lead.phone,
                "email": lead.email,
                "billing_address": ", ".join(filter(None, [primary_address.address, primary_address.city,
                                                             primary_address.pincode]))
                                   if primary_address else lead.address,
                "billing_selected_address": primary_address,
                "shipping_same_as_billing": True,
                "address": ", ".join(filter(None, [primary_address.address, primary_address.city,
                                                     primary_address.pincode]))
                           if primary_address else lead.address,
                "cake": lead.cake,
                "cake_type": lead.cake_type,
                "flavour": lead.flavour,
                "quantity": lead.quantity,
                "event_date": lead.event_date,
                "notes": "\n".join(filter(None, [lead.customer_notes, lead.cake_message, lead.remarks])),
                "selected_address": primary_address,
                "occasion": lead.occasion,
                "delivery_date": lead.delivery_date,
                "delivery_time": lead.delivery_time,
                "delivery_type": lead.delivery_type,
                "special_instructions": "\n".join(filter(None, [lead.internal_notes, lead.staff_remarks, lead.internal_sales_notes])),
                "sales_staff": lead.assigned_to,
                "total_amount": lead.estimated_amount or lead.expected_price or 0,
                "created_by": request.user,
            },
        )
        if created:
            lead_amount = Decimal(lead.estimated_amount or lead.budget or lead.expected_price or 0)
            # Only carry a rate forward when the lead actually has one; otherwise
            # the lead amount is a manually entered line amount.
            lead_rate = lead.rate
            QuotationItem.objects.create(
                quotation=quotation,
                cake=lead.cake,
                item_name=lead.cake_type or "Custom Cake",
                description="\n".join(filter(None, [lead.theme, lead.cake_message, lead.customer_notes])),
                flavour=lead.flavour,
                weight=lead.weight,
                quantity=lead.quantity or 1,
                unit_type=lead.unit_type,
                rate=lead_rate,
                manual_amount=None if lead_rate is not None else lead_amount,
                gst_percentage=0,
            )
            QuotationActivity.objects.create(
                quotation=quotation, action="Quotation Created",
                description=f"Created from {lead.lead_no}", staff=request.user,
            )
            LeadActivity.objects.create(
                lead=lead, action="Quotation Created", staff=request.user,
                description=quotation.quotation_no,
            )
            quotation.calculate_totals()
            quotation.save(update_fields=("subtotal", "gst_amount", "total_amount", "updated_on"))
            create_quotation_version(quotation, request.user)
            # The lead already synced its event; only create one here when the
            # customer was linked just now, otherwise just link the quotation.
            sync_recurring_event(quotation, create=customer_newly_linked)
        return Response({
            "message": "Quotation created." if created else "Quotation already exists.",
            "data": QuotationSerializer(quotation).data,
            "lead": LeadSerializer(lead).data,
        })

    @action(detail=True, methods=["post"], url_path="reference-images")
    def reference_images(self, request, pk=None):
        lead = self.get_object()
        uploads = request.FILES.getlist("images")
        if not uploads:
            return Response({"message": "Select at least one image."}, status=400)
        created = []
        for upload in uploads:
            try:
                ext = validate_image_upload(upload)
            except ValueError as error:
                return Response({"message": f"{upload.name}: {error}"}, status=400)
            file_name = uuid.uuid4()
            key = f"{aws_domain_name}/crm/lead/reference/{file_name}.{ext}"
            if not put_image_to_s3(upload, key):
                return Response({"message": f"{upload.name}: upload failed."}, status=500)
            created.append(LeadReferenceImage.objects.create(
                lead=lead, file_name=file_name, file_ext=ext,
                original_name=upload.name, uploaded_by=request.user,
            ))
        LeadActivity.objects.create(
            lead=lead, action="Reference Images Added", staff=request.user,
            description=f"{len(created)} image(s)",
        )
        return Response({"data": LeadReferenceImageSerializer(created, many=True).data}, status=201)

    @action(detail=True, methods=["delete"], url_path=r"reference-images/(?P<image_id>\d+)")
    def delete_reference_image(self, request, pk=None, image_id=None):
        lead = self.get_object()
        image = lead.reference_images.get(id=image_id)
        key = f"{aws_domain_name}/crm/lead/reference/{image.file_name}.{image.file_ext}"
        try:
            get_s3_session().resource("s3").Object(settings.AWS_STORAGE_BUCKET_NAME, key).delete()
        except Exception:
            pass
        image.delete()
        LeadActivity.objects.create(lead=lead, action="Reference Image Removed", staff=request.user)
        return Response(status=204)

    @action(detail=False, methods=["post"], url_path="import")
    def import_file(self, request):
        upload = request.FILES.get("file")
        if not upload:
            return Response({"message": "A CSV or XLSX file is required."}, status=400)
        ext = upload.name.rsplit(".", 1)[-1].lower()
        try:
            if ext == "csv":
                rows = list(csv.DictReader(io.StringIO(upload.read().decode("utf-8-sig"))))
            elif ext == "xlsx":
                from openpyxl import load_workbook
                sheet = load_workbook(upload, read_only=True, data_only=True).active
                values = list(sheet.values)
                headers = [str(v).strip() if v is not None else "" for v in values[0]]
                rows = [dict(zip(headers, row)) for row in values[1:]]
            else:
                return Response({"message": "Only .csv and .xlsx files are supported."}, status=400)
        except Exception:
            return Response({"message": "The uploaded file could not be read."}, status=400)

        failures, successes = [], 0
        aliases = {"Customer Name": "customer_name", "Phone": "phone", "Occasion": "occasion",
                   "Event Date": "event_date", "Cake Type": "cake_type", "Quantity": "quantity"}
        for number, raw in enumerate(rows, start=2):
            data = {aliases.get(str(k).strip(), str(k).strip().lower().replace(" ", "_")): v for k, v in raw.items()}
            if isinstance(data.get("event_date"), datetime):
                data["event_date"] = data["event_date"].date()
            data["phone"] = str(data.get("phone") or "").split(".")[0].strip()
            if Lead.objects.filter(phone=data["phone"]).exists():
                failures.append({"row": number, "reason": "Duplicate phone"})
                continue
            serializer = self.get_serializer(data=data)
            if not serializer.is_valid():
                reason = "; ".join(f"{k}: {', '.join(map(str, v))}" for k, v in serializer.errors.items())
                failures.append({"row": number, "reason": reason})
                continue
            with transaction.atomic():
                lead = serializer.save(
                    created_by=request.user,
                    customer=sync_customer_from_crm_details(serializer.validated_data, request.user),
                )
                LeadActivity.objects.create(lead=lead, action="Lead Imported", staff=request.user)
                sync_recurring_event(lead)
            successes += 1
        return Response({"success_count": successes, "failed_count": len(failures), "errors": failures})


class CRMCustomerViewSet(
    mixins.CreateModelMixin, mixins.ListModelMixin, mixins.RetrieveModelMixin, mixins.UpdateModelMixin,
    viewsets.GenericViewSet,
):
    """Safe CRM projection of authentication customers and their business history."""
    serializer_class = None

    def get_serializer_class(self):
        if self.action in {"update", "partial_update"}:
            return CRMCustomerUpdateSerializer
        return CRMCustomerUpdateSerializer

    def partial_update(self, request, *args, **kwargs):
        customer = self.get_object()
        serializer = CRMCustomerUpdateSerializer(customer, data=request.data, partial=True, context={"request": request})
        serializer.is_valid(raise_exception=True)
        serializer.save()
        refreshed = self.get_queryset().get(pk=customer.pk)
        return Response(self.summary(refreshed))

    def create(self, request, *args, **kwargs):
        serializer = CRMCustomerUpdateSerializer(data=request.data, context={"request": request})
        serializer.is_valid(raise_exception=True)
        customer = serializer.save()
        refreshed = self.get_queryset().get(pk=customer.pk)
        return Response(self.summary(refreshed), status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        return self.partial_update(request, *args, **kwargs)

    def get_queryset(self):
        Customer = get_user_model()
        not_cancelled = ~Q(crm_orders__status="Cancelled")
        queryset = Customer.objects.filter(is_staff=False)
        if getattr(self, "action", None) in {
            "restore",
            "permanent_delete",
            "deletion_dependencies_action",
            "clear_deletion_dependencies",
        }:
            pass
        elif getattr(self, "action", None) in {"retrieve", "update", "partial_update"} and is_super_admin(self.request.user):
            pass
        elif str(self.request.query_params.get("archived") or "").lower() in {"1", "true", "yes"}:
            queryset = queryset.filter(is_active=False) if is_super_admin(self.request.user) else queryset.none()
        else:
            queryset = queryset.filter(is_active=True)
        queryset = queryset.annotate(
            lead_count=Count("crm_leads", distinct=True),
            quotation_count=Count("crm_quotations", distinct=True),
            order_count=Count("crm_orders", distinct=True),
            invoice_count=Count("crm_invoices", distinct=True),
            delivery_count=Count("crm_orders__delivery", distinct=True),
            # "Valid" orders exclude Cancelled, matching CRMOrder.STATUS_CHOICES —
            # this is the figure returning-customer/revenue logic must use, not
            # the raw order_count above (kept for backward compatibility).
            valid_order_count=Count("crm_orders", filter=not_cancelled, distinct=True),
            total_revenue=models.Sum("crm_orders__total_amount", filter=not_cancelled),
            last_order_date=models.Max("crm_orders__confirmed_on", filter=not_cancelled),
        ).prefetch_related("customeraddress_set")
        term = str(self.request.query_params.get("search") or "").strip()
        if term:
            queryset = queryset.filter(
                Q(first_name__icontains=term) | Q(last_name__icontains=term) |
                Q(username__icontains=term) | Q(phone__icontains=term) |
                Q(email__icontains=term)
            )
        return queryset.order_by("-id")

    @staticmethod
    def summary(customer):
        return {
            "id": customer.id, "username": customer.username,
            "name": " ".join(filter(None, [customer.first_name, customer.last_name])).strip() or customer.username,
            "first_name": customer.first_name, "last_name": customer.last_name,
            "email": customer.email, "phone": customer.phone,
            "is_active": customer.is_active, "is_staff": customer.is_staff,
            "archived_at": customer.archived_at,
            "archived_by": customer.archived_by_id,
            "archived_by_name": str(customer.archived_by) if customer.archived_by else None,
            "date_joined": customer.date_joined, "last_login": customer.last_login,
            "phone_verified_at": customer.phone_verified_at,
            "email_verified_at": customer.email_verified_at,
            "addresses": [{"id": address.id, "title": address.title,
                           "person_name": address.person_name, "phone": address.phone,
                           "address": address.address, "city": address.city,
                           "pincode": address.pincode, "is_primary": address.is_primary,
                           "created_on": address.created_on, "updated_on": address.updated_on}
                          for address in customer.customeraddress_set.all()],
            "counts": {"leads": customer.lead_count, "quotations": customer.quotation_count,
                       "orders": customer.order_count, "invoices": customer.invoice_count,
                       "deliveries": customer.delivery_count},
            "valid_order_count": customer.valid_order_count,
            "total_revenue": customer.total_revenue or 0,
            "last_order_date": customer.last_order_date,
            "is_returning_customer": customer.valid_order_count > 1,
        }

    def destroy(self, request, *args, **kwargs):
        customer = self.get_object()
        customer.is_active = False
        customer.archived_at = timezone.now()
        customer.archived_by = request.user
        customer.save(update_fields=("is_active", "archived_at", "archived_by"))
        return Response({"message": "Customer archived."})

    @action(detail=True, methods=["post"], url_path="restore")
    def restore(self, request, pk=None):
        denied = _require_super_admin(request)
        if denied:
            return denied
        customer = self.get_object()
        customer.is_active = True
        customer.archived_at = None
        customer.archived_by = None
        customer.save(update_fields=("is_active", "archived_at", "archived_by"))
        return Response({"message": "Customer restored."})

    @staticmethod
    def deletion_dependencies(customer):
        deliveries = CRMDelivery.objects.filter(order__customer=customer).order_by("-id")
        links = [
            {
                "label": "Leads",
                "count": customer.crm_leads.count(),
                "records": [
                    {**record, "record_type": "leads"}
                    for record in customer.crm_leads.order_by("-id").values("id", "lead_no", "customer_name", "status")[:20]
                ],
            },
            {
                "label": "Quotations",
                "count": customer.crm_quotations.count(),
                "records": [
                    {**record, "record_type": "quotations"}
                    for record in customer.crm_quotations.order_by("-id").values("id", "quotation_no", "customer_name", "status")[:20]
                ],
            },
            {
                "label": "Orders",
                "count": customer.crm_orders.count(),
                "records": [
                    {**record, "record_type": "orders"}
                    for record in customer.crm_orders.order_by("-id").values("id", "order_no", "customer_name", "status")[:20]
                ],
            },
            {
                "label": "Invoices",
                "count": customer.crm_invoices.count(),
                "records": [
                    {**record, "record_type": "invoices"}
                    for record in customer.crm_invoices.order_by("-id").values("id", "invoice_no", "customer_name", "status")[:20]
                ],
            },
            {
                "label": "Deliveries",
                "count": deliveries.count(),
                "records": [
                    {
                        "id": delivery.id,
                        "delivery_no": delivery.delivery_no,
                        "customer_name": delivery.recipient_name or delivery.order.customer_name,
                        "status": delivery.delivery_status,
                        "record_type": "deliveries",
                    }
                    for delivery in deliveries[:20]
                ],
            },
        ]
        links = [link for link in links if link["count"]]
        return {"has_dependencies": bool(links), "links": links}

    @staticmethod
    def protected_error_dependencies(exc):
        protected = list(getattr(exc, "protected_objects", []) or [])
        records = []
        for obj in protected[:20]:
            record = {"id": getattr(obj, "pk", None), "message": str(obj)}
            if isinstance(obj, Lead):
                record.update({"lead_no": obj.lead_no, "customer_name": obj.customer_name, "status": obj.status, "record_type": "leads"})
            elif isinstance(obj, Quotation):
                record.update({"quotation_no": obj.quotation_no, "customer_name": obj.customer_name, "status": obj.status, "record_type": "quotations"})
            elif isinstance(obj, CRMOrder):
                record.update({"order_no": obj.order_no, "customer_name": obj.customer_name, "status": obj.status, "record_type": "orders"})
            elif isinstance(obj, CRMInvoice):
                record.update({"invoice_no": obj.invoice_no, "customer_name": obj.customer_name, "status": obj.status, "record_type": "invoices"})
            elif isinstance(obj, CRMDelivery):
                record.update({"delivery_no": obj.delivery_no, "customer_name": obj.recipient_name, "status": obj.delivery_status, "record_type": "deliveries"})
            records.append(record)
        return {
            "has_dependencies": bool(records),
            "links": [{"label": "Protected related records", "count": len(protected), "records": records}] if records else [],
        }

    @action(detail=True, methods=["get"], url_path="deletion-dependencies")
    def deletion_dependencies_action(self, request, pk=None):
        denied = _require_super_admin(request)
        if denied:
            return denied
        return Response(self.deletion_dependencies(self.get_object()))

    @action(detail=True, methods=["post"], url_path="clear-deletion-dependencies")
    def clear_deletion_dependencies(self, request, pk=None):
        denied = _require_super_admin(request)
        if denied:
            return denied
        customer = self.get_object()
        confirmation = str(request.data.get("confirmation") or "")
        if confirmation != "CLEAR CUSTOMER LINKS":
            return Response({"message": "Type CLEAR CUSTOMER LINKS to confirm clearing customer links."}, status=400)
        with transaction.atomic():
            counts = {
                "leads": customer.crm_leads.update(customer=None),
                "quotations": customer.crm_quotations.update(customer=None),
                "orders": customer.crm_orders.update(customer=None),
                "invoices": customer.crm_invoices.update(customer=None),
            }
        return Response({"message": "Customer links cleared. Business records were kept.", "counts": counts})

    @action(detail=True, methods=["delete"], url_path="permanent-delete")
    def permanent_delete(self, request, pk=None):
        denied = _require_super_admin(request)
        if denied:
            return denied
        customer = self.get_object()
        dependencies = self.deletion_dependencies(customer)
        if dependencies["has_dependencies"]:
            return Response({
                "message": "Customer cannot be permanently deleted because it has business history. Clear linked CRM records first.",
                "dependencies": dependencies,
            }, status=409)
        try:
            customer.delete()
        except ProtectedError as exc:
            return Response({
                "message": "Customer cannot be permanently deleted because it has protected related records. Clear linked CRM records first.",
                "dependencies": self.protected_error_dependencies(exc),
            }, status=409)
        return Response(status=204)

    def list(self, request, *args, **kwargs):
        return Response([self.summary(customer) for customer in self.get_queryset()])

    @staticmethod
    def _normalize_import_header(value):
        return "".join(character for character in str(value or "").casefold() if character.isalnum())

    @classmethod
    def _customer_import_field(cls, header):
        aliases = {
            "customerid": "id", "id": "id",
            "customername": "name", "name": "name", "fullname": "name",
            "firstname": "first_name", "first": "first_name",
            "lastname": "last_name", "surname": "last_name",
            "phonenumber": "phone", "mobile": "phone", "mobilenumber": "phone",
            "phone": "phone", "contact": "phone", "contactnumber": "phone",
            "emailaddress": "email", "email": "email",
            "address": "address", "primaryaddress": "address",
            "city": "city", "pincode": "pincode", "pin": "pincode",
            "postalcode": "pincode", "postcode": "pincode", "zipcode": "pincode",
            "zip": "pincode", "status": "status", "customerstatus": "status",
        }
        return aliases.get(cls._normalize_import_header(header))

    @staticmethod
    def _split_customer_name(name):
        parts = str(name or "").strip().split()
        if not parts:
            return "", ""
        return parts[0], " ".join(parts[1:])

    def _read_customer_import_rows(self, upload):
        name = str(getattr(upload, "name", "") or "")
        extension = name.rsplit(".", 1)[-1].casefold() if "." in name else ""
        if extension == "csv":
            text = upload.read().decode("utf-8-sig")
            return list(csv.reader(io.StringIO(text)))
        if extension == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet = workbook.active
            return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
        raise ValueError("Upload a CSV or XLSX customer file.")

    def _import_customer_row(self, data, request):
        customer = None
        customer_id = data.get("id")
        phone = data.get("phone")
        email = data.get("email")
        if customer_id:
            customer = get_user_model().objects.filter(id=customer_id, is_staff=False).first()
            if not customer:
                raise ValueError("Customer ID was not found.")
        if not customer and phone:
            matches = get_user_model().objects.filter(phone=phone, is_staff=False)
            if matches.count() > 1:
                raise ValueError("Multiple existing customers use this phone number.")
            customer = matches.first()
        if not customer and email:
            matches = get_user_model().objects.filter(email__iexact=email, is_staff=False)
            if matches.count() > 1:
                raise ValueError("Multiple existing customers use this email address.")
            customer = matches.first()

        name = data.get("name")
        first_name = data.get("first_name")
        last_name = data.get("last_name")
        if name and not first_name:
            first_name, split_last = self._split_customer_name(name)
            last_name = last_name or split_last
        if not str(first_name or "").strip():
            raise ValueError("Missing customer name.")
        if not str(phone or "").strip():
            raise ValueError("Missing phone number.")

        payload = {
            "first_name": first_name,
            "last_name": last_name or "",
            "email": email or "",
            "phone": phone,
            "is_active": data.get("status", "active").casefold() not in {"inactive", "archived", "disabled", "false", "0"},
        }
        address_payload = {
            "title": "Primary Address",
            "person_name": " ".join(filter(None, [payload["first_name"], payload["last_name"]])).strip(),
            "phone": phone,
            "address": data.get("address") or "",
            "city": data.get("city") or "",
            "pincode": data.get("pincode") or "",
            "is_primary": True,
        }
        if address_payload["address"] or address_payload["city"] or address_payload["pincode"]:
            primary = customer.customeraddress_set.order_by("-is_primary", "id").first() if customer else None
            if primary:
                address_payload["id"] = primary.id
            payload["addresses"] = [address_payload]
        serializer = CRMCustomerUpdateSerializer(
            customer,
            data=payload,
            partial=bool(customer),
            context={"request": request},
        )
        serializer.is_valid(raise_exception=True)
        return serializer.save(), bool(customer)

    @action(detail=False, methods=["post"], url_path="import")
    def import_file(self, request):
        upload = request.FILES.get("file")
        if not upload:
            return Response({"message": "Select a CSV or XLSX file to import."}, status=400)
        if upload.size > 5 * 1024 * 1024:
            return Response({"message": "Customer import file must be 5 MB or smaller."}, status=400)
        try:
            rows = self._read_customer_import_rows(upload)
        except Exception as exc:
            return Response({"message": str(exc)}, status=400)
        rows = [list(row) for row in rows if any(str(cell or "").strip() for cell in row)]
        if not rows:
            return Response({"message": "The import file is empty."}, status=400)
        headers = [str(cell or "").strip() for cell in rows[0]]
        mapped = [self._customer_import_field(header) for header in headers]
        unknown = [headers[index] for index, field in enumerate(mapped) if headers[index] and not field]
        if unknown:
            return Response({"message": "Some columns are not supported.", "unknown_columns": unknown}, status=400)
        if not any(field in {"name", "first_name"} for field in mapped) or "phone" not in mapped:
            return Response({"message": "Import requires customer name and phone columns."}, status=400)

        result = {"created": 0, "updated": 0, "skipped": 0, "failed": 0, "errors": [], "unknown_columns": []}
        with transaction.atomic():
            for row_number, row in enumerate(rows[1:], start=2):
                row_data = {}
                for index, field in enumerate(mapped):
                    if not field:
                        continue
                    value = row[index] if index < len(row) else ""
                    value = "" if value is None else str(value).strip()
                    if value:
                        row_data[field] = value
                if not row_data:
                    result["skipped"] += 1
                    continue
                try:
                    _, updated = self._import_customer_row(row_data, request)
                    result["updated" if updated else "created"] += 1
                except Exception as exc:
                    result["failed"] += 1
                    detail = getattr(exc, "detail", exc)
                    result["errors"].append({"row": row_number, "reason": str(detail)})
        return Response(result)

    @action(detail=False, methods=["get"], url_path="export")
    def export_file(self, request):
        queryset = self.get_queryset()
        ids = [value for value in str(request.query_params.get("ids") or "").split(",") if value.strip()]
        if ids:
            queryset = queryset.filter(id__in=ids)
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="customers.csv"'
        writer = csv.writer(response)
        writer.writerow([
            "Customer ID", "Name", "Phone", "Email", "Address", "City", "Pincode",
            "Total Orders", "Total Revenue", "Last Order", "Customer Status", "Repeat Customer",
        ])
        for customer in queryset[:1000]:
            primary = customer.customeraddress_set.order_by("-is_primary", "id").first()
            writer.writerow([
                customer.id,
                " ".join(filter(None, [customer.first_name, customer.last_name])).strip() or customer.username,
                customer.phone or "",
                customer.email or "",
                primary.address if primary else "",
                primary.city if primary else "",
                primary.pincode if primary else "",
                customer.valid_order_count,
                customer.total_revenue or 0,
                customer.last_order_date or "",
                "Active" if customer.is_active else "Inactive",
                "Yes" if customer.valid_order_count > 1 else "No",
            ])
        return response

    @action(detail=False, methods=["get"], url_path="analytics")
    def analytics(self, request):
        """Reuse the customer projection under the separately enforceable
        analytics permission consumed by the Customer Dashboard."""
        queryset = self.get_queryset()
        return Response([self.summary(customer) for customer in queryset])

    def retrieve(self, request, *args, **kwargs):
        customer = self.get_queryset().get(pk=kwargs["pk"])
        data = self.summary(customer)
        data["history"] = {
            "leads": list(customer.crm_leads.values(
                "id", "lead_no", "status", "occasion", "event_date", "follow_up_date",
                "follow_up_notes", "cake_type", "flavour", "quantity", "created_on"
            )),
            "quotations": list(customer.crm_quotations.values("id", "quotation_no", "status", "total_amount", "created_on")),
            "orders": list(customer.crm_orders.values(
                "id", "order_no", "status", "total_amount", "event_date", "confirmed_on",
                occasion=models.F("quotation__lead__occasion"),
            )),
            "invoices": list(customer.crm_invoices.values(
                "id", "invoice_no", "status", "total_amount", "amount_paid", "balance_amount", "generated_on"
            )),
            "deliveries": list(CRMDelivery.objects.filter(order__customer=customer).values(
                "id", "delivery_no", "delivery_status", "delivery_type", "pickup_time",
                "delivery_time", "porter_time", "created_on"
            )),
            "payments": list(QuotationPayment.objects.filter(quotation__customer=customer).values(
                "id", "quotation_id", "quotation__quotation_no", "amount", "payment_method", "payment_date", "created_on"
            )),
            "communications": [{
                "id": log.id, "channel": log.channel, "subject": log.subject, "message": log.message,
                "staff_name": getattr(log.staff, "first_name", None) or getattr(log.staff, "username", None),
                "created_on": log.created_on,
            } for log in customer.communication_logs.select_related("staff").all()[:100]],
        }
        data["recurring_events"] = [
            recurring_event_summary(event)
            for event in customer.recurring_events.select_related("source_lead", "source_quotation", "source_order")
        ]
        data["purchase_summary"] = self._purchase_summary(customer)
        return Response(data)

    @staticmethod
    def _purchase_summary(customer):
        """Reliable purchase-behaviour metrics. Order count/revenue/last-order
        reuse the annotations already computed once by get_queryset() (no
        extra query); only the two "favourite" breakdowns need their own
        query, and only here on the single-customer detail view (no N+1)."""
        count = customer.valid_order_count
        total_revenue = customer.total_revenue or 0
        if count == 0:
            return {
                "total_orders": 0, "total_revenue": 0, "average_order_value": 0,
                "last_order_date": None, "is_returning_customer": False,
                "favourite_occasion": None, "favourite_flavour": None,
            }
        valid_orders = customer.crm_orders.exclude(status="Cancelled")
        occasion_counts = (
            valid_orders.exclude(quotation__lead__occasion__isnull=True)
            .values("quotation__lead__occasion")
            .annotate(n=Count("id")).order_by("-n").first()
        )
        flavour_counts = (
            valid_orders.exclude(models.Q(flavour__isnull=True) | models.Q(flavour=""))
            .values("flavour").annotate(n=Count("id")).order_by("-n").first()
        )
        return {
            "total_orders": count,
            "total_revenue": total_revenue,
            "average_order_value": round(total_revenue / count, 2) if count else 0,
            "last_order_date": customer.last_order_date,
            "is_returning_customer": count > 1,
            "favourite_occasion": occasion_counts["quotation__lead__occasion"] if occasion_counts else None,
            "favourite_flavour": flavour_counts["flavour"] if flavour_counts else None,
        }

    @action(detail=False, methods=["get"], url_path="upcoming-events")
    def upcoming_events(self, request):
        """Recurring customer occasions (birthday, anniversary, ...) due within the next N days, for one-click reminders."""
        days = int(request.query_params.get("days") or 30)
        today = timezone.localdate()
        horizon = today + timedelta(days=days)
        results = []
        events = CustomerRecurringEvent.objects.select_related("customer").all()
        for event in events:
            next_occurrence = event.next_occurrence(today)
            if today <= next_occurrence <= horizon:
                summary = recurring_event_summary(event, today)
                customer = event.customer
                summary.update({
                    "customer_id": customer.id,
                    "customer_name": " ".join(filter(None, [customer.first_name, customer.last_name])).strip() or customer.username,
                    "phone": customer.phone,
                    "email": customer.email,
                })
                results.append(summary)
        results.sort(key=lambda item: item["days_until"])
        return Response(results)

    @action(detail=True, methods=["post"], url_path="send-reminder")
    def send_reminder(self, request, pk=None):
        customer = self.get_queryset().get(pk=pk)
        channel = request.data.get("channel")
        if channel not in dict(CustomerCommunicationLog.CHANNEL_CHOICES):
            return Response({"message": "A valid channel (whatsapp, sms or email) is required."}, status=400)
        event_id = request.data.get("event")
        recurring_event = CustomerRecurringEvent.objects.filter(id=event_id, customer=customer).first() if event_id else None
        if recurring_event:
            recurring_event.last_reminder_sent_on = timezone.localdate()
            recurring_event.save(update_fields=["last_reminder_sent_on", "updated_on"])
        log = CustomerCommunicationLog.objects.create(
            customer=customer, channel=channel, subject=request.data.get("subject") or "",
            message=request.data.get("message") or "", recurring_event=recurring_event,
            staff=request.user if request.user.is_authenticated else None,
        )
        return Response({"data": {"id": log.id, "channel": log.channel, "created_on": log.created_on}})


class CustomerRecurringEventViewSet(viewsets.ModelViewSet):
    """Manual CRUD for customer occasions/reminders (Customer 360 → Events tab).
    Same model the Lead pipeline already auto-populates via sync_recurring_event —
    this just exposes it as a real, staff-editable resource."""
    serializer_class = CustomerRecurringEventSerializer
    queryset = CustomerRecurringEvent.objects.select_related(
        "customer", "source_lead", "source_quotation", "source_order",
    ).all()

    def get_queryset(self):
        queryset = super().get_queryset()
        customer_id = self.request.query_params.get("customer")
        if customer_id:
            queryset = queryset.filter(customer_id=customer_id)
        return queryset

    @action(detail=True, methods=["post"])
    def complete(self, request, pk=None):
        """Marks the current reminder cycle as actioned (no message is sent —
        this only logs that staff have handled it, same semantics as the
        existing CRMCustomerViewSet.send_reminder action's date bookkeeping)."""
        event = self.get_object()
        event.last_reminder_sent_on = timezone.localdate()
        event.save(update_fields=["last_reminder_sent_on", "updated_on"])
        return Response(self.get_serializer(event).data)


def _confirm_quotation_order(quotation, user):
    """Create (or fetch) the CRMOrder backing a quotation that already has a
    recorded advance. Shared by QuotationViewSet.confirm_order and
    CRMOrderViewSet.create_direct so a directly-created order goes through the
    exact same item/activity pipeline as a normal quotation-confirmed order."""
    lead = quotation.lead
    customer_newly_linked = False
    if quotation.customer is None:
        customer = sync_customer_from_crm_details({
            "customer_name": quotation.customer_name or (lead.customer_name if lead else ""),
            "phone": quotation.phone or (lead.phone if lead else ""),
            "email": quotation.email or (lead.email if lead else ""),
            "address": quotation.address or quotation.billing_address or (lead.address if lead else ""),
        }, user)
        if customer is not None:
            quotation.customer = customer
            quotation.save(update_fields=("customer", "updated_on"))
            customer_newly_linked = True
    order, created = CRMOrder.objects.get_or_create(
        quotation=quotation,
        defaults={
            "customer": quotation.customer,
            "customer_name": quotation.customer_name or (lead.customer_name if lead else ""),
            "phone": quotation.phone or (lead.phone if lead else ""),
            "cake_type": quotation.cake_type,
            "flavour": quotation.flavour,
            "quantity": quotation.quantity or 1,
            "event_date": quotation.event_date,
            "total_amount": quotation.total_amount,
            "advance_paid": quotation.advance_amount,
            "recipient_name": quotation.customer_name,
            "recipient_phone": quotation.phone,
            "delivery_address": quotation.address,
            "delivery_date": quotation.delivery_date or (lead.delivery_date if lead else None),
            "delivery_time": quotation.delivery_time or (lead.delivery_time if lead else None),
            "delivery_method": quotation.delivery_type or (lead.delivery_type if lead else None),
            "special_instructions": quotation.special_instructions or quotation.notes,
            "assigned_staff": lead.assigned_to if lead else None,
            "confirmed_by": user,
        },
    )
    if created:
        CRMOrderItem.objects.bulk_create([
            CRMOrderItem(
                order=order, item_name=item.item_name, description=item.description,
                flavour=item.flavour, weight=item.weight, quantity=item.quantity,
                unit_type=item.unit_type,
                rate=item.rate, gst_percentage=item.gst_percentage,
                amount=item.amount, sort_order=item.sort_order,
            ) for item in quotation.items.all()
        ])
        quotation.status = "Order Confirmed"
        quotation.save(update_fields=("status", "updated_on"))
        QuotationActivity.objects.create(
            quotation=quotation, action="Order Confirmed",
            description=order.order_no, staff=user,
        )
        sync_recurring_event(order, create=customer_newly_linked)
        whatsapp_trigger_event("order_confirmed", order)
    return order, created


class QuotationViewSet(ArchivedLifecycleMixin, viewsets.ModelViewSet):
    lifecycle_label = "Quotation"
    serializer_class = QuotationSerializer
    queryset = Quotation.objects.select_related("lead", "customer", "cake").prefetch_related("versions").all()

    def get_serializer_class(self):
        return QuotationListSerializer if self.action == "list" else QuotationSerializer

    def permanent_delete_blocker(self, instance):
        if CRMOrder.objects.filter(quotation=instance).exists():
            return "Quotation cannot be permanently deleted because it has a related order. Archive it instead."
        return None

    def prepare_permanent_delete(self, instance):
        instance.payments.all().delete()
        instance.versions.all().delete()

    def deletion_dependencies(self, instance):
        links = []
        orders = CRMOrder.objects.filter(quotation=instance).order_by("-id")
        if orders.exists():
            links.append({
                "label": "Orders",
                "count": orders.count(),
                "records": [{
                    "id": order.id,
                    "order_no": order.order_no,
                    "customer_name": order.customer_name,
                    "status": order.status,
                    "record_type": "orders",
                } for order in orders[:20]],
            })
        return {"has_dependencies": bool(links), "links": links}

    def perform_create(self, serializer):
        customer = serializer.validated_data.get("customer")
        if customer is None:
            customer = sync_customer_from_crm_details(serializer.validated_data, self.request.user)
        quotation = serializer.save(
            created_by=self.request.user,
            customer=customer,
        )
        QuotationActivity.objects.create(
            quotation=quotation, action="Quotation Created",
            description="Manual quotation", staff=self.request.user,
        )
        create_quotation_version(quotation, self.request.user)
        sync_recurring_event(quotation)

    def perform_update(self, serializer):
        previous = serializer.instance.status
        previous_event = _event_fields(serializer.instance, "occasion", "event_date", "customer_id")
        customer = serializer.validated_data.get("customer", serializer.instance.customer)
        if customer is None and any(field in serializer.validated_data for field in ("customer_name", "phone", "email", "address", "billing_address")):
            customer = sync_customer_from_crm_details(serializer.validated_data, self.request.user)
        quotation = serializer.save(customer=customer) if customer is not None else serializer.save()
        if _event_fields(quotation, "occasion", "event_date", "customer_id") != previous_event:
            sync_recurring_event(quotation)
        if previous != quotation.status and quotation.lead:
            LeadActivity.objects.create(
                lead=quotation.lead, action=f"Quotation {quotation.status}",
                staff=self.request.user, description=quotation.quotation_no,
            )
        if previous != quotation.status:
            QuotationActivity.objects.create(
                quotation=quotation, action=f"Quotation {quotation.status}",
                staff=self.request.user,
            )
            if quotation.status == "Sent":
                whatsapp_trigger_event("quotation_sent", quotation)

    @action(detail=True, methods=["post"], url_path="generate-version")
    def generate_version(self, request, pk=None):
        """Validate/save the edited quotation and explicitly freeze a new document version."""
        with transaction.atomic():
            quotation = Quotation.objects.select_for_update().get(pk=self.get_object().pk)
            previous_event = _event_fields(quotation, "occasion", "event_date", "customer_id")
            serializer = self.get_serializer(quotation, data=request.data, partial=True)
            serializer.is_valid(raise_exception=True)
            quotation = serializer.save()
            if _event_fields(quotation, "occasion", "event_date", "customer_id") != previous_event:
                sync_recurring_event(quotation)
            version = create_quotation_version(quotation, request.user)
            QuotationActivity.objects.create(
                quotation=quotation, action="Quotation Generated",
                description=f"Version {version.version_number}", staff=request.user,
            )
        return Response(self.get_serializer(quotation).data)

    @action(detail=False, methods=["get"])
    def pending(self, request):
        leads = Lead.objects.filter(status="Won", quotation__isnull=True).select_related("assigned_to")
        return Response(LeadSerializer(leads, many=True).data)

    @action(detail=False, methods=["get"], url_path="customer-search")
    def customer_search(self, request):
        term = str(request.query_params.get("q") or "").strip()
        if len(term) < 2:
            return Response([])
        cache_version = cache.get("crm:customer-search:version", 1)
        cache_key = f"crm:customer-search:customers-only:{cache_version}:{term.casefold()}"
        cached = cache.get(cache_key)
        if cached is not None:
            return Response(cached)
        Customer = get_user_model()
        customers = Customer.objects.filter(is_staff=False)
        token_query = Q()
        for token in [part for part in term.split() if part]:
            token_query &= (
                Q(first_name__icontains=token) | Q(last_name__icontains=token) |
                Q(phone__icontains=token) | Q(email__icontains=token)
            )
        customers = customers.filter(
            Q(first_name__icontains=term) | Q(last_name__icontains=term) |
            Q(phone__icontains=term) | Q(email__icontains=term) | token_query
        ).prefetch_related("customeraddress_set")[:20]
        results = [{
            "id": customer.id,
            "name": " ".join(filter(None, [customer.first_name, customer.last_name])).strip() or customer.username,
            "phone": customer.phone or customer.username,
            "email": customer.email,
            "is_active": customer.is_active,
            "customer_since": customer.date_joined,
            "addresses": [{"id": address.id, "title": address.title,
                           "person_name": address.person_name, "phone": address.phone,
                           "address": address.address,
                           "city": address.city, "pincode": address.pincode,
                           "is_primary": address.is_primary}
                          for address in customer.customeraddress_set.all()],
            "history": {
                "quotations": customer.crm_quotations.count(),
                "orders": customer.crm_orders.count(),
                "invoices": customer.crm_invoices.count(),
                "deliveries": CRMDelivery.objects.filter(order__customer=customer).count(),
            },
        } for customer in customers]
        normalized = term.casefold()
        digits = "".join(filter(str.isdigit, term))
        results.sort(key=lambda item: 0 if (
            (digits and "".join(filter(str.isdigit, str(item["phone"] or ""))) == digits) or
            str(item["email"] or "").casefold() == normalized or
            str(item["name"] or "").casefold() == normalized
        ) else 1)
        cache.set(cache_key, results, 300)
        return Response(results)

    @action(detail=False, methods=["post"], url_path="save-customer-address")
    def save_customer_address(self, request):
        from customer.models import CustomerAddress
        customer_id = request.data.get("customer")
        if not customer_id or not str(request.data.get("address") or "").strip():
            return Response({"message": "Customer and address are required."}, status=400)
        address = CustomerAddress.objects.create(
            customer_id=customer_id, title=request.data.get("title") or "Delivery Address",
            person_name=request.data.get("person_name"), phone=request.data.get("phone"),
            address=str(request.data["address"]).strip(), city=request.data.get("city"),
            pincode=request.data.get("pincode"), is_primary=bool(request.data.get("is_primary")),
            created_by=request.user,
        )
        version_key = "crm:customer-search:version"
        try:
            cache.incr(version_key)
        except ValueError:
            cache.set(version_key, 2, None)
        return Response({"id": address.id, "title": address.title,
                         "person_name": address.person_name, "phone": address.phone,
                         "address": address.address,
                         "city": address.city, "pincode": address.pincode,
                         "is_primary": address.is_primary}, status=201)

    @action(detail=True, methods=["post"], url_path="workflow-action")
    def workflow_action(self, request, pk=None):
        quotation = self.get_object()
        transitions = {
            "send": ({"Draft"}, "Sent", "Quotation Sent"),
            "accept": ({"Sent"}, "Accepted", "Quotation Accepted"),
            "cancel": ({"Draft", "Sent", "Accepted"}, "Cancelled", "Quotation Cancelled"),
        }
        action_name = request.data.get("action")
        if action_name not in transitions:
            return Response({"message": "Unknown quotation action."}, status=400)
        if is_quotation_locked(quotation):
            return Response(
                {"message": "This quotation is finalized with a confirmed order and cannot be changed."},
                status=400,
            )
        allowed, new_status, activity = transitions[action_name]
        if quotation.status not in allowed:
            return Response(
                {"message": f"This action is not available while the quotation is {quotation.status}."},
                status=400,
            )
        quotation.status = new_status
        quotation.save(update_fields=("status", "updated_on"))
        QuotationActivity.objects.create(
            quotation=quotation, action=activity, staff=request.user,
        )
        if quotation.lead:
            LeadActivity.objects.create(
                lead=quotation.lead, action=activity, staff=request.user,
                description=quotation.quotation_no,
            )
        return Response(self.get_serializer(quotation).data)

    @action(detail=True, methods=["post"], url_path="record-payment")
    def record_payment(self, request, pk=None):
        with transaction.atomic():
            quotation = Quotation.objects.select_for_update().get(pk=self.get_object().pk)
            if quotation.status != "Accepted":
                message = "Advance has already been recorded." if quotation.advance_received or quotation.advance_amount > 0 else "Accept the quotation before recording an advance."
                return Response({"message": message}, status=400)
            paid = sum(quotation.payments.values_list("amount", flat=True))
            if quotation.advance_received or quotation.advance_amount > 0 or paid > 0:
                return Response({"message": "Advance has already been recorded."}, status=400)
            serializer = QuotationPaymentSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            amount = serializer.validated_data["amount"]
            if amount <= 0:
                return Response({"amount": ["Amount must be greater than zero."]}, status=400)
            if amount > quotation.total_amount:
                return Response({"amount": ["Payment cannot exceed the remaining balance."]}, status=400)
            payment = serializer.save(quotation=quotation, collected_by=request.user)
            quotation.advance_received = True
            quotation.advance_amount = payment.amount
            quotation.payment_date = payment.payment_date
            quotation.payment_method = payment.payment_method
            quotation.transaction_reference = payment.transaction_reference
            quotation.status = "Advance Received"
            quotation.save()
            QuotationActivity.objects.create(
                quotation=quotation, action="Advance Payment Recorded",
                staff=request.user, description=str(payment.amount),
            )
        return Response(self.get_serializer(quotation).data, status=201)

    @action(detail=True, methods=["post"], url_path="confirm-order")
    def confirm_order(self, request, pk=None):
        with transaction.atomic():
            # Lock only the Quotation row. PostgreSQL cannot apply FOR UPDATE to the
            # nullable side of the OUTER JOIN produced by select_related("lead").
            quotation = Quotation.objects.select_for_update().get(pk=self.get_object().pk)
            if quotation.status not in {"Advance Received", "Order Confirmed"}:
                return Response(
                    {"message": "Record an advance payment before confirming the order."}, status=400
                )
            if not quotation.advance_received or quotation.advance_amount <= 0:
                return Response({"message": "A valid advance payment is required."}, status=400)
            order, created = _confirm_quotation_order(quotation, request.user)
        return Response(
            {"message": "Order confirmed." if created else "Order already exists.",
             "data": CRMOrderSerializer(order).data},
            status=201 if created else 200,
        )


class CRMOrderViewSet(ArchivedLifecycleMixin, viewsets.ModelViewSet):
    lifecycle_label = "Order"
    serializer_class = CRMOrderSerializer
    queryset = CRMOrder.objects.select_related(
        "quotation", "quotation__lead", "customer", "delivery_staff",
        "assigned_staff", "assigned_baker",
    ).prefetch_related("items", "quotation__activities", "quotation__lead__reference_images").all()

    def get_queryset(self):
        queryset = super().get_queryset()
        start, end = self.request.query_params.get("from_date"), self.request.query_params.get("to_date")
        if start:
            queryset = queryset.filter(Q(delivery_date__gte=start) | Q(delivery_date__isnull=True, event_date__gte=start))
        if end:
            queryset = queryset.filter(Q(delivery_date__lte=end) | Q(delivery_date__isnull=True, event_date__lte=end))
        return queryset

    def permanent_delete_blocker(self, instance):
        if CRMInvoice.objects.filter(order=instance).exists() or CRMDelivery.objects.filter(order=instance).exists():
            return "Order cannot be permanently deleted because it has invoice or delivery history. Archive it instead."
        return None

    def deletion_dependencies(self, instance):
        links = []
        invoices = CRMInvoice.objects.filter(order=instance).order_by("-id")
        if invoices.exists():
            links.append({
                "label": "Invoices",
                "count": invoices.count(),
                "records": [{
                    "id": invoice.id,
                    "invoice_no": invoice.invoice_no,
                    "customer_name": invoice.customer_name,
                    "status": invoice.status,
                    "record_type": "invoices",
                } for invoice in invoices[:20]],
            })
        deliveries = CRMDelivery.objects.filter(order=instance).order_by("-id")
        if deliveries.exists():
            links.append({
                "label": "Deliveries",
                "count": deliveries.count(),
                "records": [{
                    "id": delivery.id,
                    "delivery_no": delivery.delivery_no,
                    "customer_name": delivery.recipient_name or instance.customer_name,
                    "status": delivery.delivery_status,
                    "record_type": "deliveries",
                } for delivery in deliveries[:20]],
            })
        return {"has_dependencies": bool(links), "links": links}

    def create(self, request, *args, **kwargs):
        return Response(
            {"message": "Orders are created by confirming a quotation."}, status=405
        )

    @action(detail=False, methods=["post"], url_path="direct")
    def create_direct(self, request):
        """Create an order directly (walk-in/phone customer) without a Lead or
        Quotation in the UI. Internally builds a minimal manual Quotation and
        runs it through the same order-creation pipeline as confirm_order, so
        the resulting order is indistinguishable from a normal one."""
        try:
            advance_value = Decimal(str(request.data.get("advance_amount") or "0"))
        except (InvalidOperation, TypeError):
            return Response({"advance_amount": ["Enter a valid advance amount."]}, status=400)
        if advance_value <= 0:
            return Response({"advance_amount": ["A valid advance payment is required to create a direct order."]}, status=400)
        # Validate the advance payment details (method/date/reference/remarks,
        # collected the same way QuotationViewSet.record_payment does) before
        # creating anything, so a bad payment form doesn't leave a stray
        # quotation/order behind.
        payment_fields = {"payment_date", "payment_method", "transaction_reference", "customer_remarks"}
        payment_data = {"amount": str(advance_value), **{key: request.data.get(key) for key in payment_fields if request.data.get(key) not in (None, "")}}
        payment_serializer = QuotationPaymentSerializer(data=payment_data)
        payment_serializer.is_valid(raise_exception=True)
        # Create the quotation first (with its items) so total_amount is
        # computed before validating the advance against it — validate()
        # compares advance_amount to a live-computed total that only reflects
        # top-level rate/quantity at validation time, not nested items.
        payload = {key: value for key, value in request.data.items() if key not in {"advance_amount", "advance_received", "status", *payment_fields}}
        create_serializer = QuotationSerializer(data=payload)
        create_serializer.is_valid(raise_exception=True)
        with transaction.atomic():
            customer = create_serializer.validated_data.get("customer")
            if customer is None:
                customer = sync_customer_from_crm_details(create_serializer.validated_data, request.user)
            quotation = create_serializer.save(created_by=request.user, customer=customer)
            sync_recurring_event(quotation)
            confirm_serializer = QuotationSerializer(
                quotation,
                data={"status": "Advance Received", "advance_received": True, "advance_amount": str(advance_value)},
                partial=True,
            )
            confirm_serializer.is_valid(raise_exception=True)
            quotation = confirm_serializer.save()
            payment_serializer.save(quotation=quotation, collected_by=request.user)
            order, created = _confirm_quotation_order(quotation, request.user)
        return Response(
            {"message": "Order created." if created else "Order already exists.",
             "data": CRMOrderSerializer(order).data},
            status=201 if created else 200,
        )

    @action(detail=True, methods=["post"], url_path="record-payment")
    def record_payment(self, request, pk=None):
        with transaction.atomic():
            order = CRMOrder.objects.select_for_update().select_related("quotation").get(pk=pk)
            invoice = None
            if hasattr(order, "invoice"):
                invoice = CRMInvoice.objects.select_for_update().get(order=order)
                if invoice.status == "Cancelled":
                    return Response({"message": "Payments cannot be recorded against a cancelled invoice."}, status=400)
            if order.status == "Cancelled":
                return Response({"message": "Payments cannot be recorded against a cancelled order."}, status=400)
            serializer = QuotationPaymentSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            amount = serializer.validated_data["amount"]
            if amount <= 0 or amount > order.balance_amount:
                return Response(
                    {"amount": ["Enter an amount greater than zero and no more than the balance due."]},
                    status=400,
                )
            reference = serializer.validated_data.get("transaction_reference")
            if reference and QuotationPayment.objects.filter(
                quotation=order.quotation, transaction_reference=reference,
            ).exists():
                return Response({"transaction_reference": ["This payment reference has already been recorded."]}, status=400)
            payment = serializer.save(
                quotation=order.quotation, collected_by=request.user
            )
            order.advance_paid += payment.amount
            order.save(update_fields=("advance_paid", "updated_on"))
            if invoice:
                invoice.amount_paid = order.advance_paid
                invoice.balance_amount = order.balance_amount
                invoice.status = "Paid" if order.balance_amount == 0 else invoice.status
                invoice.save(update_fields=("amount_paid", "balance_amount", "status", "updated_on"))
            QuotationActivity.objects.create(
                quotation=order.quotation, action="Order Payment Recorded",
                description=str(payment.amount), staff=request.user,
            )
        return Response(self.get_serializer(order).data, status=201)

    @action(detail=True, methods=["post"], url_path="workflow-action")
    def workflow_action(self, request, pk=None):
        order = self.get_object()
        action_name = request.data.get("action")
        if action_name == "cancel_order":
            uncancellable = {"Delivery", "Out for Delivery", "Delivered", "Completed", "Closed", "Cancelled"}
            if order.status in uncancellable:
                return Response(
                    {"message": f"This order cannot be cancelled while it is {order.status}."}, status=400
                )
            reason = str(request.data.get("cancellation_reason") or "").strip()
            if not reason:
                return Response(
                    {"message": "A cancellation reason is required.", "fields": ["cancellation_reason"]},
                    status=400,
                )
            order.status = "Cancelled"
            order.cancellation_reason = reason
            order.cancelled_at = timezone.now()
            order.cancelled_by = request.user
            order.save(update_fields=("status", "cancellation_reason", "cancelled_at", "cancelled_by", "updated_on"))
            QuotationActivity.objects.create(
                quotation=order.quotation, action="Order Cancelled",
                description=reason, staff=request.user,
            )
            return Response(self.get_serializer(order).data)
        transitions = {
            "start_processing": ({"Confirmed", "Production"}, "Processing"),
            "mark_ready": ({"Processing"}, "Ready"),
            "dispatch_delivery": ({"Invoice Generated", "Out for Delivery"}, "Delivery"),
            "complete_delivery": ({"Delivery", "Out for Delivery", "Delivered"}, "Completed"),
        }
        if action_name not in transitions:
            return Response({"message": "Unknown order action."}, status=400)
        expected, new_status = transitions[action_name]
        if order.status not in expected:
            return Response(
                {"message": f"This action is not available while the order is {order.status}."},
                status=400,
            )
        if action_name == "dispatch_delivery":
            required = ("recipient_name", "recipient_phone", "delivery_address",
                        "delivery_date", "delivery_time", "delivery_method")
            missing = [field for field in required if not getattr(order, field)]
            if missing:
                return Response(
                    {"message": "Complete all required delivery information first.",
                     "fields": missing}, status=400,
                )
        if action_name == "complete_delivery" and order.balance_amount > 0:
            return Response(
                {"message": "Collect the remaining balance before completing this order."},
                status=400,
            )
        order.status = new_status
        order.save(update_fields=("status", "updated_on"))
        QuotationActivity.objects.create(
            quotation=order.quotation, action=f"Order {new_status}",
            description=order.order_no, staff=request.user,
        )
        return Response(self.get_serializer(order).data)

    @action(detail=True, methods=["post"], url_path="generate-invoice")
    def generate_invoice(self, request, pk=None):
        with transaction.atomic():
            # Lock the Order without joining its nullable Customer. PostgreSQL
            # rejects FOR UPDATE when it is applied across the nullable side of
            # an outer join. The customer is loaded separately only if needed.
            order = CRMOrder.objects.select_for_update().select_related("quotation").get(pk=pk)
            if order.status not in {"Ready", "Invoice Pending", "Invoice Generated"}:
                return Response({"message": "Mark the order as Ready before generating its invoice."}, status=400)
            quotation = order.quotation
            invoice, created = CRMInvoice.objects.get_or_create(order=order, defaults={
                "customer": order.customer, "customer_name": order.customer_name, "phone": order.phone,
                "address": order.delivery_address or quotation.address or quotation.billing_address,
                "subtotal": quotation.subtotal, "gst_amount": quotation.gst_amount,
                "total_amount": order.total_amount, "amount_paid": order.advance_paid,
                "balance_amount": order.balance_amount,
                "status": "Paid" if order.balance_amount <= 0 else "Generated",
                "generated_by": request.user,
            })
            if created:
                order.status = "Invoice Pending"
                order.save(update_fields=("status", "updated_on"))
                QuotationActivity.objects.create(quotation=quotation, action="Invoice Generated", description=invoice.invoice_no, staff=request.user)
                whatsapp_trigger_event("invoice_generated", invoice)
        return Response({"message": "Invoice generated." if created else "Invoice already exists.", "data": CRMInvoiceSerializer(invoice).data}, status=201 if created else 200)


class CRMInvoiceViewSet(ArchivedLifecycleMixin, viewsets.ReadOnlyModelViewSet):
    lifecycle_label = "Invoice"
    serializer_class = CRMInvoiceSerializer
    queryset = CRMInvoice.objects.select_related("order", "order__quotation", "customer").prefetch_related(
        "order__items", "order__quotation__payments", "order__quotation__activities",
    ).all()

    def get_queryset(self):
        queryset = super().get_queryset()
        start, end = self.request.query_params.get("from_date"), self.request.query_params.get("to_date")
        if start:
            queryset = queryset.filter(generated_on__date__gte=start)
        if end:
            queryset = queryset.filter(generated_on__date__lte=end)
        return queryset

    def permanent_delete_blocker(self, instance):
        return None

    def deletion_dependencies(self, instance):
        return {"has_dependencies": False, "links": []}

    @action(detail=True, methods=["post"], url_path="cancel")
    def cancel(self, request, pk=None):
        with transaction.atomic():
            invoice = CRMInvoice.objects.select_for_update().select_related("order", "order__quotation").get(pk=pk)
            order = CRMOrder.objects.select_for_update().get(pk=invoice.order_id)
            delivery = None
            if hasattr(order, "delivery"):
                delivery = CRMDelivery.objects.select_for_update().get(order=order)
                if delivery.delivery_status != "Pending":
                    return Response(
                        {"message": f"This invoice cannot be cancelled while delivery is {delivery.delivery_status}."},
                        status=400,
                    )
            if invoice.status not in {"Generated", "Issued"}:
                return Response({"message": f"This invoice cannot be cancelled while it is {invoice.status}."}, status=400)
            invoice.status = "Cancelled"
            invoice.save(update_fields=("status", "updated_on"))
            if order.status in {"Invoice Pending", "Invoice Generated"}:
                order.status = "Ready"
                order.save(update_fields=("status", "updated_on"))
            if delivery:
                delivery.delivery_status = "Cancelled"
                delivery.save(update_fields=("delivery_status", "updated_on"))
            QuotationActivity.objects.create(
                quotation=invoice.order.quotation, action="Invoice Cancelled",
                description=invoice.invoice_no, staff=request.user,
            )
        return Response(self.get_serializer(invoice).data)

    @action(detail=True, methods=["post"], url_path="record-payment")
    def record_payment(self, request, pk=None):
        with transaction.atomic():
            invoice = CRMInvoice.objects.select_for_update().select_related("order", "order__quotation").get(pk=pk)
            order = CRMOrder.objects.select_for_update().get(pk=invoice.order_id)
            if invoice.status == "Cancelled":
                return Response({"message": "Payments cannot be recorded against a cancelled invoice."}, status=400)
            if order.status == "Cancelled":
                return Response({"message": "Payments cannot be recorded against a cancelled order."}, status=400)
            serializer = QuotationPaymentSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            amount = serializer.validated_data["amount"]
            if amount <= 0 or amount > invoice.balance_amount:
                return Response({"amount": ["Enter an amount greater than zero and no more than the balance due."]}, status=400)
            reference = serializer.validated_data.get("transaction_reference")
            if reference and QuotationPayment.objects.filter(
                quotation=invoice.order.quotation, transaction_reference=reference,
            ).exists():
                return Response({"transaction_reference": ["This payment reference has already been recorded."]}, status=400)
            payment = serializer.save(quotation=invoice.order.quotation, collected_by=request.user)
            order.advance_paid += payment.amount
            order.save(update_fields=("advance_paid", "updated_on"))
            invoice.amount_paid = order.advance_paid
            invoice.balance_amount = order.balance_amount
            if invoice.balance_amount <= 0:
                invoice.status = "Paid"
            invoice.save(update_fields=("amount_paid", "balance_amount", "status", "updated_on"))
            QuotationActivity.objects.create(quotation=invoice.order.quotation, action="Invoice Payment Recorded", description=str(payment.amount), staff=request.user)
            if invoice.status == "Paid":
                QuotationActivity.objects.create(quotation=invoice.order.quotation, action="Payment Completed", description=invoice.invoice_no, staff=request.user)
        return Response(self.get_serializer(invoice).data, status=201)

    @action(detail=True, methods=["post"], url_path="payment-completed")
    def payment_completed(self, request, pk=None):
        invoice = self.get_object()
        if invoice.status == "Cancelled":
            return Response({"message": "Cancelled invoices cannot be marked paid."}, status=400)
        if invoice.balance_amount > 0:
            return Response({"message": "Collect the full balance before completing payment."}, status=400)
        if invoice.status != "Paid":
            invoice.status = "Paid"
            invoice.save(update_fields=("status", "updated_on"))
            QuotationActivity.objects.create(quotation=invoice.order.quotation, action="Payment Completed", description=invoice.invoice_no, staff=request.user)
        return Response(self.get_serializer(invoice).data)

    @action(detail=True, methods=["post"], url_path="move-to-delivery")
    def move_to_delivery(self, request, pk=None):
        with transaction.atomic():
            invoice = CRMInvoice.objects.select_for_update().select_related("order", "order__quotation").get(pk=pk)
            if invoice.status != "Paid" or invoice.balance_amount > 0:
                return Response({"message": "Complete payment before moving to Delivery."}, status=400)
            order = invoice.order
            delivery_type = "Customer Pickup" if str(order.delivery_method or "").lower() == "pickup" else "Magilshe Delivery"
            scheduled = None
            if order.delivery_date:
                from datetime import datetime, time as time_cls
                naive = datetime.combine(order.delivery_date, order.delivery_time or time_cls(0, 0))
                scheduled = timezone.make_aware(naive) if timezone.is_naive(naive) else naive
            defaults = {
                "delivery_type": delivery_type,
                "recipient_name": order.recipient_name or order.customer_name,
                "recipient_phone": order.recipient_phone or order.phone,
                "created_by": request.user,
            }
            if delivery_type == "Customer Pickup":
                defaults.update({"pickup_time": scheduled, "pickup_notes": order.delivery_notes})
                if order.delivery_address:
                    defaults["pickup_location"] = order.delivery_address
            else:
                defaults.update({
                    "delivery_location": order.delivery_address, "delivery_time": scheduled,
                    "delivery_notes": order.delivery_notes,
                })
            delivery, created = CRMDelivery.objects.get_or_create(order=order, defaults=defaults)
            order.status = "Delivery"
            order.save(update_fields=("status", "updated_on"))
            if created:
                QuotationActivity.objects.create(quotation=order.quotation, action="Moved to Delivery", description=delivery.delivery_no, staff=request.user)
        return Response({"message": "Delivery created." if created else "Delivery already exists.", "data": CRMDeliverySerializer(delivery).data}, status=201 if created else 200)


def _missing_delivery_fields(delivery):
    """Fields required to move a delivery from Pending to Out for Delivery, per delivery_type."""
    if delivery.delivery_type == "Customer Pickup":
        checks = [
            ("pickup_person_type", delivery.pickup_person_type),
            ("pickup_person_name", delivery.pickup_person_name),
            ("pickup_contact", delivery.pickup_contact),
            ("pickup_time", delivery.pickup_time),
            ("pickup_location", delivery.pickup_location),
        ]
        if delivery.pickup_person_type == "Other":
            checks.append(("pickup_person_relationship", delivery.pickup_person_relationship))
    elif delivery.delivery_type == "Porter Delivery":
        checks = [
            ("external_platform", delivery.external_platform_id),
            ("porter_name", delivery.porter_name),
            ("porter_contact", delivery.porter_contact),
            ("porter_location", delivery.porter_location),
            ("porter_time", delivery.porter_time),
        ]
    else:
        checks = [
            ("delivery_person", delivery.delivery_person_id),
            ("delivery_time", delivery.delivery_time),
            ("delivery_location", delivery.delivery_location),
        ]
    return [key for key, value in checks if not value]


class CRMDeliveryViewSet(ArchivedLifecycleMixin, viewsets.ModelViewSet):
    lifecycle_label = "Delivery"
    serializer_class = CRMDeliverySerializer
    queryset = CRMDelivery.objects.select_related("order", "order__quotation", "order__invoice", "order__customer", "customer_address", "delivery_person", "external_platform", "external_delivery_person").prefetch_related(
        "order__items", "order__quotation__activities", "order__customer__customeraddress_set", "damage_images",
    ).all()

    def get_queryset(self):
        queryset = super().get_queryset()
        start, end = self.request.query_params.get("from_date"), self.request.query_params.get("to_date")
        scheduled = Q()
        if start:
            scheduled &= (Q(pickup_time__date__gte=start) | Q(delivery_time__date__gte=start) | Q(porter_time__date__gte=start))
        if end:
            scheduled &= (Q(pickup_time__date__lte=end) | Q(delivery_time__date__lte=end) | Q(porter_time__date__lte=end))
        return queryset.filter(scheduled) if start or end else queryset

    def permanent_delete_blocker(self, instance):
        return None

    def deletion_dependencies(self, instance):
        return {"has_dependencies": False, "links": []}

    def create(self, request, *args, **kwargs):
        return Response({"message": "Deliveries are created from paid invoices."}, status=405)

    @action(detail=True, methods=["post"], url_path="cancel")
    def cancel(self, request, pk=None):
        delivery = self.get_object()
        if delivery.delivery_status != "Pending":
            return Response({"message": f"This delivery cannot be cancelled while it is {delivery.delivery_status}."}, status=400)
        delivery.delivery_status = "Cancelled"
        delivery.save(update_fields=("delivery_status", "updated_on"))
        QuotationActivity.objects.create(
            quotation=delivery.order.quotation, action="Delivery Cancelled",
            description=delivery.delivery_no, staff=request.user,
        )
        return Response(self.get_serializer(delivery).data)

    @action(detail=True, methods=["post"], url_path="workflow-action")
    def workflow_action(self, request, pk=None):
        with transaction.atomic():
            # Keep nullable/reverse relations outside the locking query for the
            # same PostgreSQL FOR UPDATE constraint. Order and Quotation are
            # required relations and remain safe inner joins.
            delivery = CRMDelivery.objects.select_for_update().select_related(
                "order", "order__quotation"
            ).get(pk=pk)
            action_name = request.data.get("action")
            if action_name == "start" and delivery.delivery_status == "Pending":
                missing = _missing_delivery_fields(delivery)
                if missing:
                    return Response({"message": "Complete the required delivery details first.", "fields": missing}, status=400)
                delivery.delivery_status = "Out for Delivery"
                activity = "Delivery Started"
                update_fields = ("delivery_status", "updated_on")
            elif action_name == "deliver" and delivery.delivery_status == "Out for Delivery":
                delivery.delivery_status = "Delivered Successfully"
                delivery.completed_at = timezone.now()
                activity = "Delivered Successfully"
                update_fields = ("delivery_status", "completed_at", "updated_on")
            elif action_name == "mark_damaged" and delivery.delivery_status == "Out for Delivery":
                description = str(request.data.get("damage_description") or "").strip()
                if not description:
                    return Response({"message": "Damage description is required."}, status=400)
                delivery.delivery_status = "Damaged During Delivery"
                delivery.damage_description = description
                activity = "Damaged During Delivery"
                update_fields = ("delivery_status", "damage_description", "updated_on")
            elif action_name == "mark_failed" and delivery.delivery_status == "Out for Delivery":
                reason = str(request.data.get("failure_reason") or "").strip()
                if not reason:
                    return Response({"message": "Failure reason is required."}, status=400)
                delivery.delivery_status = "Delivery Failed"
                delivery.failure_reason = reason
                activity = "Delivery Failed"
                update_fields = ("delivery_status", "failure_reason", "updated_on")
            else:
                return Response({"message": f"This action is not available while delivery is {delivery.delivery_status}."}, status=400)
            delivery.save(update_fields=update_fields)
            order = delivery.order
            QuotationActivity.objects.create(quotation=order.quotation, action=activity, description=delivery.delivery_no, staff=request.user)
            if delivery.delivery_status == "Out for Delivery":
                whatsapp_trigger_event("out_for_delivery", delivery)
            if delivery.delivery_status == "Delivered Successfully":
                whatsapp_trigger_event("delivered", delivery)
                order.status = "Closed"
                order.save(update_fields=("status", "updated_on"))
                invoice = order.invoice
                invoice.status = "Paid"
                invoice.save(update_fields=("status", "updated_on"))
                quotation = order.quotation
                quotation.status = "Closed"
                quotation.save(update_fields=("status", "updated_on"))
                if quotation.lead:
                    quotation.lead.status = "Completed"
                    quotation.lead.save(update_fields=("status", "updated_on"))
                    LeadActivity.objects.create(lead=quotation.lead, action="Workflow Closed", description=order.order_no, staff=request.user)
                QuotationActivity.objects.create(quotation=quotation, action="Workflow Closed", description=order.order_no, staff=request.user)
        return Response(self.get_serializer(delivery).data)

    @action(detail=True, methods=["post"], url_path="damage-images")
    def damage_images(self, request, pk=None):
        delivery = self.get_object()
        uploads = request.FILES.getlist("images")
        if not uploads:
            return Response({"message": "Select at least one image."}, status=400)
        created = []
        for upload in uploads:
            try:
                ext = validate_image_upload(upload)
            except ValueError as error:
                return Response({"message": f"{upload.name}: {error}"}, status=400)
            file_name = uuid.uuid4()
            key = f"{aws_domain_name}/crm/delivery/damage/{file_name}.{ext}"
            if not put_image_to_s3(upload, key):
                return Response({"message": f"{upload.name}: upload failed."}, status=500)
            created.append(CRMDeliveryDamageImage.objects.create(
                delivery=delivery, file_name=file_name, file_ext=ext,
                original_name=upload.name, uploaded_by=request.user,
            ))
        return Response({"data": CRMDeliveryDamageImageSerializer(created, many=True).data}, status=201)

    @action(detail=True, methods=["delete"], url_path=r"damage-images/(?P<image_id>\d+)")
    def delete_damage_image(self, request, pk=None, image_id=None):
        delivery = self.get_object()
        image = delivery.damage_images.get(id=image_id)
        key = f"{aws_domain_name}/crm/delivery/damage/{image.file_name}.{image.file_ext}"
        try:
            get_s3_session().resource("s3").Object(settings.AWS_STORAGE_BUCKET_NAME, key).delete()
        except Exception:
            pass
        image.delete()
        return Response(status=204)


class CRMOptionViewSet(viewsets.ModelViewSet):
    serializer_class = CRMOptionSerializer
    queryset = CRMOption.objects.all()

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.request.query_params.get("type"):
            queryset = queryset.filter(option_type=self.request.query_params["type"])
        if self.request.query_params.get("active") in {"true", "1"}:
            queryset = queryset.filter(is_active=True)
        return queryset


class CanViewAgenda(BasePermission):
    """The agenda is shown on the CRM Calendar, the Customers dashboard and the
    Customer 360 profile, so any of those pages' view permissions grants it."""
    message = "You do not have permission to view the CRM calendar."
    codes = {"crm.dashboard.view", "customers.analytics.view", "customers.data.view"}

    def has_permission(self, request, view):
        return is_super_admin(request.user) or bool(self.codes & effective_permissions(request.user))


class CRMAgendaView(APIView):
    """GET crm/agenda/?start=YYYY-MM-DD&end=YYYY-MM-DD[&customer=<id>]
    Day-wise orders, deliveries, pending payments, follow-ups, customer
    events/reminders and completed work (see crm/agenda.py)."""
    permission_classes = (IsAuthenticated, CanViewAgenda)

    def get(self, request):
        today = timezone.localdate()
        try:
            start = datetime.strptime(request.query_params.get("start") or today.isoformat(), "%Y-%m-%d").date()
            end = datetime.strptime(
                request.query_params.get("end") or (start + timedelta(days=6)).isoformat(), "%Y-%m-%d",
            ).date()
        except ValueError:
            return Response({"message": "start and end must be dates (YYYY-MM-DD)."}, status=400)
        if end < start:
            return Response({"message": "end must be on or after start."}, status=400)
        if (end - start).days >= MAX_RANGE_DAYS:
            return Response({"message": f"The range can span at most {MAX_RANGE_DAYS} days."}, status=400)
        customer = request.query_params.get("customer")
        if customer and not str(customer).isdigit():
            return Response({"message": "customer must be a customer id."}, status=400)
        return Response(build_agenda(start, end, customer_id=int(customer) if customer else None, today=today))
